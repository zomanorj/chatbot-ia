import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

import docx
import openpyxl
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_TEXT_CHARS = 1_000_000

CONTENT_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".md": "text/markdown",
}
SUPPORTED_EXTENSIONS = tuple(CONTENT_TYPES)


@dataclass(frozen=True)
class Block:
    """A unit of text. Standalone blocks (table rows, FAQ sections) are never merged."""

    text: str
    standalone: bool = False


class ExtractionError(Exception):
    """A document problem the owner can fix. `code` is shown to them, translated."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


def _zip_contains(data: bytes, member: str) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return member in archive.namelist()
    except zipfile.BadZipFile:
        return False


def detect_content_type(filename: str, data: bytes) -> str:
    """Trust the extension only when the bytes agree with it."""
    extension = PurePath(filename).suffix.lower()
    content_type = CONTENT_TYPES.get(extension)
    if content_type is None:
        raise ExtractionError("unsupported_type")
    valid = {
        ".pdf": lambda: data.startswith(b"%PDF-"),
        ".docx": lambda: _zip_contains(data, "word/document.xml"),
        ".xlsx": lambda: _zip_contains(data, "xl/workbook.xml"),
    }.get(extension, lambda: b"\x00" not in data[:8192])
    if not valid():
        raise ExtractionError("corrupted_file")
    return content_type


def _clean(text: str) -> str:
    return re.sub(r"[ \t\u00a0]+", " ", text).strip()


def _paragraphs(text: str) -> list[str]:
    blocks = re.split(r"\n\s*\n", text.replace("\r\n", "\n"))
    return [cleaned for block in blocks if (cleaned := _clean(re.sub(r"\s*\n\s*", " ", block)))]


def _text(data: bytes) -> list[Block]:
    return [Block(p) for p in _paragraphs(_decode(data))]


def _markdown(data: bytes) -> list[Block]:
    """One standalone block per heading section, so each FAQ entry is retrieved on its own."""
    sections: list[list[str]] = [[]]
    for paragraph in _paragraphs(_decode(data)):
        if paragraph.startswith("#") and sections[-1]:
            sections.append([])
        sections[-1].append(
            paragraph.lstrip("# ").strip() if paragraph.startswith("#") else paragraph
        )
    has_headings = len(sections) > 1
    return [Block("\n".join(section), standalone=has_headings) for section in sections if section]


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("unreadable_text")


def _format_row(headers: list[str], values: list[object], prefix: str = "") -> str | None:
    pairs = [
        f"{header}: {_clean(str(value))}" if header else _clean(str(value))
        for header, value in zip(headers, values, strict=False)
        if value is not None and str(value).strip()
    ]
    return prefix + "; ".join(pairs) if pairs else None


def _pdf(data: bytes) -> list[Block]:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("encrypted_pdf")
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except PdfReadError as exc:
        raise ExtractionError("corrupted_file") from exc
    blocks = [Block(p) for p in _paragraphs(text)]
    if not blocks:
        # Scanned documents contain images only; OCR is not supported yet.
        raise ExtractionError("no_text")
    return blocks


def _docx(data: bytes) -> list[Block]:
    document = docx.Document(io.BytesIO(data))
    blocks = [Block(_clean(p.text)) for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        headers = [_clean(h) for h in rows[0]]
        blocks += [row for values in rows[1:] if (row := _row(headers, values))]
    return blocks


def _row(headers: list[str], values: list[object], prefix: str = "") -> Block | None:
    text = _format_row(headers, values, prefix)
    return Block(text, standalone=True) if text else None


def _xlsx(data: bytes) -> list[Block]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks: list[Block] = []
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True) if any(row)]
        if not rows:
            continue
        headers = [_clean(str(h)) if h is not None else "" for h in rows[0]]
        prefix = f"{sheet.title} — " if len(workbook.worksheets) > 1 else ""
        blocks += [row for values in rows[1:] if (row := _row(headers, values, prefix))]
    workbook.close()
    return blocks


def _csv(data: bytes) -> list[Block]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = [
        row for row in csv.reader(io.StringIO(text), dialect) if any(cell.strip() for cell in row)
    ]
    if not rows:
        return []
    headers = [_clean(h) for h in rows[0]]
    return [row for values in rows[1:] if (row := _row(headers, list(values)))]


EXTRACTORS = {
    "application/pdf": _pdf,
    CONTENT_TYPES[".docx"]: _docx,
    CONTENT_TYPES[".xlsx"]: _xlsx,
    "text/csv": _csv,
    "text/plain": _text,
    "text/markdown": _markdown,
}


def extract_blocks(content_type: str, data: bytes) -> list[Block]:
    """Return text blocks: paragraphs, table rows or document sections."""
    try:
        blocks = EXTRACTORS[content_type](data)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("corrupted_file") from exc
    if sum(len(b.text) for b in blocks) > MAX_TEXT_CHARS:
        raise ExtractionError("too_much_text")
    if not blocks:
        raise ExtractionError("no_text")
    return blocks
